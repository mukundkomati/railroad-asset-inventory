

import re
import openpyxl
import pandas as pd

SRC = "Data_Engineer_TechnicalChallenge.xlsx"
SHEET = "TS Program Est"
COST_TYPE = "Construction w/ Materials"


def read_rates(ws):
    return {
        "track_miles_m1": ws["I15"].value,
        "track_miles_m2": ws["I16"].value,
        "ts_per_track_mile": ws["I18"].value,
        "per_turnout": ws["I23"].value,
        "industry_reference": ws["I24"].value,
        "per_curve": ws["I27"].value,
    }


def read_rom(ws):
    years = {c: ws.cell(4, c).value.replace("State ", "").strip()
             for c in range(9, 15)}

    rom = {}
    for r in range(5, 9):
        name = ws.cell(r, 8).value
        if not name:
            continue
        row = {years[c]: ws.cell(r, c).value
               for c in range(9, 15) if ws.cell(r, c).value}
        row["total"] = ws.cell(r, 15).value
        rom[name.strip()] = row          # a couple have trailing spaces
    return rom


def curves(ws, rates):
    out = []
    for r in range(5, 13):
        label = ws.cell(r, 1).value
        if not label:
            continue

        mp = float(re.search(r"MP\s*([\d.]+)", label).group(1))
        track = "M2" if label.strip().endswith("M2") else "M1"

        out.append({
            "program": "Curve Rail Replacement",
            "asset_name": f"Curve MP {mp} {track}",
            "asset_kind": "Curve",
            "milepost_start": mp, "milepost_end": mp, "track": track,
            "quantity": 1, "quantity_uom": "curve", "is_industry": 0,
            "amount": rates["per_curve"],
            "fiscal_year": ws.cell(r, 3).value.strip(),
            "allocation_method": "Unit rate I27",
            "source_ref": f"{SHEET}!A{r}",
        })
    return out


def timber_and_surfacing(rates):
    out = []
    for track, miles, fy in [("M1", rates["track_miles_m1"], "FY30"),
                             ("M2", rates["track_miles_m2"], "FY31")]:
        out.append({
            "program": "Timber & Surfacing",
            "asset_name": f"T&S Segment {track} MP 15.0-32.5",
            "asset_kind": "Track Segment",
            "milepost_start": 15.0, "milepost_end": 32.5, "track": track,
            "quantity": miles, "quantity_uom": "track miles", "is_industry": 0,
            "amount": miles * rates["ts_per_track_mile"],
            "fiscal_year": fy,
            "allocation_method": "Unit rate I18; mainline-to-year split assumed",
            "source_ref": f"{SHEET}!E5",
        })
    return out


def turnouts(rom, rates):
    out = []
    for _, t in pd.read_csv("lookups/turnouts.csv").iterrows():
        out.append({
            "program": "Turnout Replacement",
            "asset_name": f"{t['name']} TO MP {t.milepost}",
            "asset_kind": "Industry Turnout" if t.is_industry else "Turnout",
            "milepost_start": t.milepost, "milepost_end": t.milepost,
            "track": t.track,
            "quantity": t.quantity, "quantity_uom": "turnouts",
            "is_industry": t.is_industry,
            "amount": t.quantity * rates["per_turnout"],
            "fiscal_year": t.assigned_fy,
            "allocation_method": f"Unit rate I23; FY basis: {t.fy_basis}",
            "source_ref": f"{SHEET}!{t.source_row}",
        })

    leftover = rom["TO Replacement"]["total"] - 11 * rates["per_turnout"]
    if leftover:
        out.append({
            "program": "Turnout Replacement",
            "asset_name": "Unknown Turnout",
            "asset_kind": "Turnout",
            "milepost_start": None, "milepost_end": None, "track": None,
            "quantity": None, "quantity_uom": None, "is_industry": 0,
            "amount": leftover,
            "fiscal_year": "FY27",
            "allocation_method": "Unattributed - I8 has no turnout to attach to",
            "source_ref": f"{SHEET}!I8",
        })
    return out


def stations(rom):
    lookup = pd.read_csv("lookups/stations.csv")
    each = rom["Stations Tie Replacement"]["total"] / len(lookup)

    out = []
    for _, s in lookup.iterrows():
        out.append({
            "program": "Stations Tie Replacement",
            "asset_name": f"{s['name']} Tie Replacement",
            "asset_kind": "Station Track",
            "milepost_start": s.milepost, "milepost_end": s.milepost,
            "track": s.track,
            "quantity": 1, "quantity_uom": "station", "is_industry": 0,
            "amount": each,
            "fiscal_year": s.assigned_fy,
            "allocation_method": "Even split of I7 - no unit rate exists",
            "source_ref": f"{SHEET}!{s.source_row}",
        })
    return out


def extract():
    ws = openpyxl.load_workbook(SRC, data_only=True)[SHEET]
    rates = read_rates(ws)
    rom = read_rom(ws)

    rows = (curves(ws, rates)
            + timber_and_surfacing(rates)
            + turnouts(rom, rates)
            + stations(rom))

    df = pd.DataFrame(rows)
    df["cost_type"] = COST_TYPE
    return df, rates, rom


def check(df, rates, rom):
    out = []

    for prog, rom_name in [("Curve Rail Replacement", "Curve Replacement"),
                           ("Timber & Surfacing", "T&S"),
                           ("Stations Tie Replacement", "Stations Tie Replacement"),
                           ("Turnout Replacement", "TO Replacement")]:
        got = df.loc[df.program == prog, "amount"].sum()
        out.append((f"{prog} ties", got == rom[rom_name]["total"]))

    out.append(("total 12,487,000", df.amount.sum() == 12_487_000))
    out.append(("11 turnouts",
                df.loc[df.program == "Turnout Replacement", "quantity"].sum() == 11))

    # I24 divided by I23 is what tells us four of the turnouts are Industry
    industry = df.loc[df.is_industry == 1, "quantity"].sum()
    out.append(("industry count = I24/I23",
                industry == rates["industry_reference"] / rates["per_turnout"]))

    # Year by year, not just the total. Put a turnout in the wrong year and the
    # grand total still ties while FY28 and FY29 quietly go wrong.
    by_year = {}
    for prog in rom.values():
        for fy, amount in prog.items():
            if fy != "total":
                by_year[fy] = by_year.get(fy, 0) + amount

    mine = df.groupby("fiscal_year").amount.sum()
    for fy in sorted(by_year):
        out.append((f"{fy} ties", mine.get(fy, 0) == by_year[fy]))

    return out


if __name__ == "__main__":
    df, rates, rom = extract()

    for name, passed in check(df, rates, rom):
        print(f"  {'ok  ' if passed else 'FAIL'} {name}")

    print()
    print(df.groupby("program").agg(assets=("asset_name", "size"),
                                    units=("quantity", "sum"),
                                    cost=("amount", "sum")).to_string())
    print()
    print(df.groupby("fiscal_year").amount.sum().to_string())

    df.to_csv("raw_ts_program.csv", index=False)
    print(f"\n{len(df)} rows -> raw_ts_program.csv")