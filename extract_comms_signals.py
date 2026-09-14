"""Tab 1 -> raw_comms_signals.csv"""

import openpyxl
import pandas as pd

SRC = "Data_Engineer_TechnicalChallenge.xlsx"
SHEET = "Comms and Signals"

BLOCK_START = {1: "A", 2: "G", 3: "M"}
FIRST_ROW, LAST_ROW = 2, 24


def extract():
    ws = openpyxl.load_workbook(SRC, data_only=True)[SHEET]
    rows = []

    for phase, start in BLOCK_START.items():
        c = openpyxl.utils.column_index_from_string(start)

        for r in range(FIRST_ROW, LAST_ROW + 1):
            # the column headed "Phase N" holds the cost type, not the phase
            cost_type = ws.cell(r, c).value
            if not cost_type:
                continue                    # separator row

            rows.append({
                "phase": phase,
                "cost_type": cost_type.strip(),
                "milepost": ws.cell(r, c + 1).value,
                "location_name": ws.cell(r, c + 2).value.strip(),
                "estimated_cost": ws.cell(r, c + 3).value,
                "quarter": ws.cell(r, c + 4).value.strip(),
                "calendar_year": ws.cell(r, c + 5).value,
                "source_ref": f"{SHEET}!{chr(ord(start) + 3)}{r}",
            })

    return pd.DataFrame(rows)


def check(df):
    by_type = df.groupby("cost_type").estimated_cost.sum()
    material_per_phase = (df[df.cost_type == "Material"]
                          .groupby("phase").estimated_cost.sum())

    return [
        ("54 cost cells", len(df) == 54),
        ("18 assets", df.location_name.nunique() == 18),
        ("3 cost types each", (df.groupby("location_name").size() == 3).all()),
        ("total 11,229,000", df.estimated_cost.sum() == 11_229_000),
        ("design 876,000", by_type["Design"] == 876_000),
        ("material 3,300,000", by_type["Material"] == 3_300_000),
        ("construction 7,053,000", by_type["Construction"] == 7_053_000),
        ("material 1.1m per phase", (material_per_phase == 1_100_000).all()),
        ("no nulls", not df.isna().any().any()),
    ]


if __name__ == "__main__":
    df = extract()

    for name, passed in check(df):
        print(f"  {'ok  ' if passed else 'FAIL'} {name}")

    print()
    print(df.pivot_table(index="phase", columns="cost_type",
                         values="estimated_cost", aggfunc="sum",
                         margins=True).to_string())

    df.to_csv("raw_comms_signals.csv", index=False)
    print(f"\n{len(df)} rows -> raw_comms_signals.csv")